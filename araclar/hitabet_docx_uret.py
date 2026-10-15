#!/usr/bin/env python3
"""
hitabet_docx_uret.py
Converts speech markdown files directly into clean, standard Word documents (.docx)
without extraneous presentation design or callout boxes:
- docs/planli_eskitme_hitabet_metni_en.md -> docs/planli_eskitme_hitabet_metni_en.docx
- docs/planli_eskitme_hitabet_metni_tr.md -> docs/planli_eskitme_hitabet_metni_tr.docx
"""

import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def md_to_docx(md_path, docx_path):
    if not os.path.exists(md_path):
        return
        
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()
    
    # Standard margins (1 inch)
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    blocks = content.split('\n\n')
    for b in blocks:
        b = b.strip()
        if not b or b == '---':
            continue
            
        # Markdown H1 (# ...)
        if b.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(b[2:].strip())
            r.font.name = 'Calibri'
            r.font.size = Pt(14)
            r.font.bold = True
            continue
            
        # Major headings (MODEL ..., THE ROOT ..., CONCLUSION., etc.)
        if (b.startswith('MODEL ') or b.startswith('THE ROOT ') or b.startswith('CONCLUSION.') or 
            b.startswith('THE JURISTIC ') or b.startswith('NOW THE DEFENSES') or 
            b.startswith('KÖK TAVSİF') or b.startswith('NETİCE.') or b.startswith('ÜRETİCİ SAVUNMALARI')):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(b.strip())
            r.font.name = 'Calibri'
            r.font.size = Pt(13)
            r.font.bold = True
            continue
            
        # Evidence headings
        if (b.startswith('OUR ') and 'EVIDENCE:' in b) or 'DELİLİMİZ:' in b:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(b.strip())
            r.font.name = 'Calibri'
            r.font.size = Pt(11.5)
            r.font.bold = True
            continue
            
        # Lines in block
        lines = b.split('\n')
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue
                
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            # Bullets and indents
            if line_str.startswith('* ') or line_str.startswith('- '):
                p.paragraph_format.left_indent = Inches(0.25)
                line_str = '•  ' + line_str[2:].strip()
            elif line_str.startswith('   ') or line_str.startswith('  '):
                p.paragraph_format.left_indent = Inches(0.4)
                
            # Parse inline bold and italic
            tokens = []
            pattern = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|[^*]+)')
            for m in pattern.finditer(line_str):
                chunk = m.group(0)
                if chunk.startswith('**') and chunk.endswith('**'):
                    tokens.append((chunk[2:-2], True, False))
                elif chunk.startswith('*') and chunk.endswith('*'):
                    tokens.append((chunk[1:-1], False, True))
                else:
                    tokens.append((chunk, False, False))
                    
            for text_part, is_b, is_i in tokens:
                run = p.add_run(text_part)
                run.font.name = 'Calibri'
                run.font.size = Pt(11)
                run.font.bold = is_b
                run.font.italic = is_i
                
    doc.save(docx_path)
    print(f'Saved: {docx_path}')

if __name__ == '__main__':
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    md_to_docx(os.path.join(base_dir, 'docs', 'planli_eskitme_hitabet_metni_en.md'),
               os.path.join(base_dir, 'docs', 'planli_eskitme_hitabet_metni_en.docx'))
    md_to_docx(os.path.join(base_dir, 'docs', 'planli_eskitme_hitabet_metni_tr.md'),
               os.path.join(base_dir, 'docs', 'planli_eskitme_hitabet_metni_tr.docx'))
