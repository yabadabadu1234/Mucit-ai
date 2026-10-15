#!/usr/bin/env python3
"""
sunum_docx_uret.py
Generates the presentation Word document (sunum.docx) by embedding
high-resolution screenshots of the presentation HTML (docs/planli_eskitme_sunum_panosu.html)
page by page in landscape orientation with full dark aesthetic and exact fit.
"""

import os
import struct
from docx import Document
from docx.shared import Inches, Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def get_png_dimensions(file_path):
    """Read width and height of a PNG file from its header without external heavy libraries."""
    with open(file_path, 'rb') as f:
        data = f.read(24)
        if data[:8] == b'\x89PNG\r\n\x1a\n' and data[12:16] == b'IHDR':
            width, height = struct.unpack('>LL', data[16:24])
            return width, height
    raise ValueError(f"Not a valid PNG file: {file_path}")

def generate_sunum_docx(output_paths=None):
    if output_paths is None:
        output_paths = [
            os.path.abspath("sunum.docx"),
            os.path.abspath("docs/sunum.docx")
        ]

    slides_dir = os.path.abspath("docs/sunum_gorseller")
    slide_files = [f"slide_0{i}.png" for i in range(8)]

    for s in slide_files:
        full_p = os.path.join(slides_dir, s)
        if not os.path.exists(full_p):
            raise FileNotFoundError(f"Slide screenshot not found: {full_p}. Run araclar/ekran_goruntusu_al.cjs first.")

    doc = Document()

    # Document properties
    props = doc.core_properties
    props.title = "Planlı Eskitme Hukuki Modeli - Aksiyomatik Sunum Panosu"
    props.subject = "Planlı Eskitme ve Platform Tekellerinin Formel Fıkhî & Hukuki Çözümlemesi"
    props.author = "Mucit-AI"
    props.category = "Sunum Belgesi"
    props.comments = "HTML sunum panosunun (planli_eskitme_sunum_panosu.html) yüksek çözünürlüklü slayt dökümü."

    # Dark background color for Word (#07090E matching presentation)
    background = OxmlElement('w:background')
    background.set(qn('w:color'), '07090E')
    doc.element.insert(0, background)

    # Word settings to display background shape
    settings = doc.settings.element
    display_bg = OxmlElement('w:displayBackgroundShape')
    settings.append(display_bg)

    # Section / Page configuration (A4 Landscape)
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width = Mm(297)   # 11.693 inches
    sec.page_height = Mm(210)  # 8.268 inches

    margin_in = 0.35
    sec.top_margin = Inches(margin_in)
    sec.bottom_margin = Inches(margin_in)
    sec.left_margin = Inches(margin_in)
    sec.right_margin = Inches(margin_in)

    # Usable printable area with safety margin
    max_print_w = (297 / 25.4) - (2 * margin_in) - 0.25   # ~10.74 in
    max_print_h = (210 / 25.4) - (2 * margin_in) - 0.25   # ~7.32 in

    print(f"Generating presentation with {len(slide_files)} slides...")
    print(f"Max printable width: {max_print_w:.2f} in, Max printable height: {max_print_h:.2f} in")

    for idx, slide_name in enumerate(slide_files):
        img_path = os.path.join(slides_dir, slide_name)
        px_w, px_h = get_png_dimensions(img_path)
        aspect = px_w / px_h

        # Proportional fit
        if (max_print_h * aspect) <= max_print_w:
            render_h = max_print_h
            render_w = render_h * aspect
        else:
            render_w = max_print_w
            render_h = render_w / aspect

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(0)

        run = p.add_run()
        run.font.size = Pt(1)
        run.add_picture(img_path, width=Inches(render_w), height=Inches(render_h))

        # Inline page break inside the same paragraph (no empty paragraph created!)
        if idx < len(slide_files) - 1:
            run.add_break(WD_BREAK.PAGE)

        print(f"Slide {idx} ({slide_name}): {px_w}x{px_h} -> rendered as {render_w:.2f} x {render_h:.2f} in")

    for out_p in output_paths:
        out_dir = os.path.dirname(out_p)
        if out_dir and not os.path.exists(out_dir):
            os.makedirs(out_dir, exist_ok=True)
        doc.save(out_p)
        print(f"Saved: {out_p} ({os.path.getsize(out_p):,} bytes)")

if __name__ == "__main__":
    generate_sunum_docx()
