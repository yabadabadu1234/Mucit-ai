import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter as L
from model import *
S=json.load(open('sonuc.json'))
wb=Workbook()
b=Font(bold=True); hf=PatternFill('solid',fgColor='1F3A5F'); hw=Font(bold=True,color='FFFFFF'); inp=PatternFill('solid',fgColor='FFF3D6')
# --- Girdiler
g=wb.active; g.title='Girdiler'
g['A1']='MİKRO MÜKELLEFİYET MALİ ETKİ MODELİ: GİRDİLER (sarı hücreler değiştirilebilir)'; g['A1'].font=Font(bold=True,size=13)
g['A2']='Aktif senaryo (1=Kötümser, 2=Orta, 3=İyimser):'; g['D2']=2; g['D2'].fill=inp
g['A3']='Tüm tutarlar sabit 2026 TL. Girdilerin çoğu ölçülmemiş varsayımdır (Kaynak sütununa bakınız).'
satir=[('Brüt asgari ücret (TL/ay)','W',33030,33030,33030,'Birden çok kaynakta tutarlı; resmî metin okunmadı'),
('Ortalama yıllık hasılat (asgari ücret katı; veri: girişim ciro/firma, 2024)','ort_H_W',43.0,43.0,43.0,'Veri: 2024 genç/mikro ciro ortalaması (ciro/firma, veri_kalibrasyon.json)'),
('Hasılat dağılımı yayılımı (σ, log-normal)','sigma',2.0,1.75,1.5,'Varsayım: 1,25–2,0 aralığı; veriyle belirlenmez'),
('Katılımcı sayısı (kararlı durum, kişi)','n_toplam',119000,238000,475000,'Kötümser: F 3-I 268; İyimser: F 3-I 269 oranları × ILO 2025 + F 3-I 273 gayrıresmî tanımı (n_veri.py); Orta: geometrik ortalama (Claude)'),
('Halen kayıtlı olup rejime geçenlerin payı (s)','s_kayitli',0.30,0.15,0.05,'Varsayım; prim isteğe bağlı olduğundan sonucu çok etkiler'),
('Gelire bağlı prim: tam prim eşiği (brüt asgari ücret katı)','prim_k',6,6,6,'Teklif Madde 5 (o); genç simülasyonunda seçildi (prim_esigi_out2.txt); F 3-I 304'),
('Tam prim tutarı (TL/ay)','prim_ay',11808.25,11808.25,11808.25,'5510 m.81: 33.030 x %35,75'),
('Gelire bağlı primin tahsil oranı','prim_toplama',1,1,1,'Tevkifatla birlikte kaynakta kesilir (tasarım); ölçülmemiş'),
('Eski kayıtlıların sabit primi sürdürme payı (0 = hiç sürdürmez)','prim_devam',0,0,0,'Sabit prim kalktığı için 0'),
('Maaşlı çalışanın mikro mükellef olma payı (u)','u_ucretli',0.25,0.20,0.15,'Varsayım; maaşlı isterse girer (F 3-I 304)'),
('Tevkifat oranı (gayrisafi hasılatın)','duz_oran',0.135,0.135,0.135,'Teklif Madde 1 (e): %13,5 ile %27 arası; F 3-I 299'),
('BES aktarım oranı (tevkifatın)','bes',0.03,0.03,0.03,'Teklif Madde 9'),
('Devlet katkısı oranı','devlet_katki',0.20,0.20,0.20,'4632 Ek 1 + 10811 sayılı Karar (ikincil)'),
('BES vazgeçme payı','bes_vazgecme',0.20,0.35,0.50,'Varsayım'),
('MYÖ prim oranı','myo_oran',0.21,0.21,0.21,'5510 m.81'),
('Hibe prim süresi (ay)','myo_gun_ay',3,3,3,'Teklif 5510 Ek 25 (90 gün)'),
('90 gün primi hak edenlerin payı (çakışmasız)','myo_uygun',0.5,0.35,0.2,'Varsayım'),
('Donanım alan payı','kdv_donanim_pay',0.5,0.4,0.3,'Varsayım'),
('Yıllık donanım KDV ortalaması (TL)','kdv_donanim_kdv',5000,5000,5000,'Varsayım'),
('KDV mahsup üst sınırı (tevkifatın oranı)','kdv_mahsup_tavan',0.5,0.5,0.5,'Teklif Madde 3 önerisi'),
('Eski rejimde kâr marjı','marj',0.5,0.5,0.5,'Varsayım'),
('Eski 4/b yıllık asgari prim (TL)','prim_eski',141699,141699,141699,'33.030×%35,75×12'),
('Eski prim ödeme oranı','prim_odeme',0.8,0.8,0.8,'Varsayım'),
('Ücretlinin marjinal oranı','r_marj_ucretli',0.27,0.27,0.27,'Varsayım'),
('Ücretlinin eski rejimde uyum oranı','uyum_ucretli',0.5,0.5,0.5,'Varsayım'),
('Kurumsal alıcı payı','kurumsal_pay',0.7,0.7,0.7,'Varsayım'),
('Kurumlar vergisi oranı','kv',0.25,0.25,0.25,'Varsayım (kanunda %25)'),
('Sahte fatura payı (kurumsal akışta, φ)','phi',0.02,0.01,0.005,'Varsayım'),
('Önlem etkinliği','onlem',0.3,0.5,0.7,'Varsayım'),
('KGF kullanım oranı','kgf_alim',0.08,0.05,0.03,'Varsayım'),
('KGF temerrüt oranı','kgf_temerrut',0.15,0.10,0.06,'Varsayım'),
('KGF kurtarma oranı','kgf_kurtarma',0.3,0.3,0.3,'Varsayım'),
('Kefalet tutarı (hasılatın payı)','kgf_pay',0.5,0.5,0.5,'Teklif Madde 8'),
('PTT UETS birim gideri (TL)','ptt_birim',50,50,50,'Varsayım'),
('UETS tahsis payı','ptt_pay',0.8,0.8,0.8,'Varsayım'),
('Tek seferlik BT gideri (TL)','bt_tek',250e6,250e6,250e6,'Varsayım'),
('Yıllık BT işletme gideri (TL)','bt_yillik',50e6,50e6,50e6,'Varsayım'),
('İtiraz oranı','itiraz_oran',0.04,0.04,0.04,'Varsayım'),
('İtiraz birim gideri (TL)','itiraz_maliyet',2000,2000,2000,'Varsayım'),
('Denetim gideri (brüt tevkifatın oranı)','denetim',0.01,0.01,0.01,'Varsayım')]
for j,h in enumerate(['Girdi','Kod','Kötümser','Orta','İyimser','AKTİF','Kaynak/durum']):
    c=g.cell(row=5,column=1+j,value=h); c.font=hw; c.fill=hf
ADR={}
for i,(ad,kod,a,o,iy,k) in enumerate(satir):
    r=6+i; g.cell(row=r,column=1,value=ad); g.cell(row=r,column=2,value=kod)
    for j,v in enumerate((a,o,iy)): c=g.cell(row=r,column=3+j,value=v); c.fill=inp
    g.cell(row=r,column=6,value=f"=INDEX(C{r}:E{r},$D$2)").font=b; g.cell(row=r,column=7,value=k); ADR[kod]=f"Girdiler!$F${r}"
g.column_dimensions['A'].width=52; g.column_dimensions['B'].width=16; g.column_dimensions['G'].width=46
# tarife
t0=6+len(satir)+2
g.cell(row=t0,column=1,value='GVK m.103 tarifesi (2026 parantez içi tutarlar)').font=b
for j,h in enumerate(['Dilim üst sınırı (TL)','Oran','Artış (Δ)','Alt sınır']): g.cell(row=t0+1,column=1+j,value=h).font=b
rows=[(190000,0.15),(400000,0.20),(1000000,0.27),(5300000,0.35),(1e12,0.40)]
for i,(e,r_) in enumerate(rows):
    r=t0+2+i; g.cell(row=r,column=1,value=e); g.cell(row=r,column=2,value=r_)
    g.cell(row=r,column=3,value=(f"=B{r}" if i==0 else f"=B{r}-B{r-1}")); g.cell(row=r,column=4,value=(0 if i==0 else f"=A{r-1}"))
TR=(t0+2,t0+6)  # rows
alt=f"Girdiler!$D${TR[0]}:$D${TR[1]}"; dl=f"Girdiler!$C${TR[0]}:$C${TR[1]}"
# --- Dağılım ve hesap
d=wb.create_sheet('Hesap'); K=60
heads=['k','H alt','H üst','H orta','ağırlık','Tevkifat matrahı (gayrisafi)','Tevkifat T','Etkin oran','BES+Devlet katkısı','90 gün prim','KDV mahsubu','Eski vergi (kayıtlı)','Eski prim','Eski vergi (ücretli)','Gelire bağlı yeni prim']
for j,h in enumerate(heads): c=d.cell(row=1,column=1+j,value=h); c.font=hw; c.fill=hf
W_=ADR['W']; lo=f"(0.1*{W_})"; hi=f"(24*{W_})"
mu=f"LN({ADR['ort_H_W']}*{W_})-{ADR['sigma']}^2/2"; sg=ADR['sigma']
for i in range(K):
    r=2+i; d.cell(row=r,column=1,value=i)
    d.cell(row=r,column=2,value=f"={lo}*({hi}/{lo})^(A{r}/{K})"); d.cell(row=r,column=3,value=f"={lo}*({hi}/{lo})^((A{r}+1)/{K})"); d.cell(row=r,column=4,value=f"=SQRT(B{r}*C{r})")
    d.cell(row=r,column=5,value=f"=(NORMSDIST((LN(C{r})-{mu})/{sg})-NORMSDIST((LN(B{r})-{mu})/{sg}))/(NORMSDIST((LN({hi})-{mu})/{sg})-NORMSDIST((LN({lo})-{mu})/{sg}))")
    d.cell(row=r,column=6,value=f"=D{r}")
    d.cell(row=r,column=7,value=f"={ADR['duz_oran']}*F{r}")
    d.cell(row=r,column=8,value=f"=G{r}/D{r}")
    d.cell(row=r,column=9,value=f"={ADR['bes']}*G{r}*(1-{ADR['bes_vazgecme']})*(1+{ADR['devlet_katki']})")
    d.cell(row=r,column=10,value=f"=IF(D{r}>=12*{W_},{ADR['myo_uygun']}*{ADR['myo_gun_ay']}*{W_}*{ADR['myo_oran']},0)")
    d.cell(row=r,column=11,value=f"={ADR['kdv_donanim_pay']}*MIN({ADR['kdv_mahsup_tavan']}*G{r},{ADR['kdv_donanim_kdv']})")
    d.cell(row=r,column=12,value=f"=SUMPRODUCT(({ADR['marj']}*D{r}>{alt})*({ADR['marj']}*D{r}-{alt})*{dl})")
    d.cell(row=r,column=13,value=f"={ADR['prim_odeme']}*{ADR['prim_eski']}*(1-{ADR['prim_devam']})")
    d.cell(row=r,column=14,value=f"={ADR['uyum_ucretli']}*{ADR['r_marj_ucretli']}*{ADR['marj']}*D{r}")
    d.cell(row=r,column=15,value=f"=12*{ADR['prim_ay']}*{ADR['prim_toplama']}*MIN(MAX(((D{r}-G{r})/12-{W_}/3)/({ADR['prim_k']}*{W_}-{W_}/3),0),1)")
last=1+K
# --- Sonuç
s=wb.create_sheet('Sonuc'); s['A1']='KARARLI DURUM (4. YIL) SONUÇLARI: aktif senaryo'; s['A1'].font=Font(bold=True,size=13)
def E(col): return f"SUMPRODUCT(Hesap!$E$2:$E${last},Hesap!${col}$2:${col}${last})"
N=ADR['n_toplam']; sK=ADR['s_kayitli']; u=ADR['u_ucretli']
SR=[('Ort. hasılat (TL)',f"={E('D')}"),('Ort. tevkifat (TL)',f"={E('G')}"),('Etkin tevkifat oranı',"=B3/B2"),
('Katılımcı (kişi)',f"={N}"),('Brüt tevkifat (TL)',"=B5*B3"),('BES + Devlet katkısı (TL)',f"=B5*{E('I')}"),('90 gün prim hibesi (TL)',f"=B5*{E('J')}"),('Donanım KDV mahsubu (TL)',f"=B5*{E('K')}"),
('Yutulma: vergi (TL)',f"=B5*({sK}*{E('L')}+{u}*{E('N')})"),('Yutulma: prim (TL)',f"=B5*{sK}*{E('M')}"),
('Sahte fatura kaybı (TL)',f"=MAX({ADR['phi']}*(1-{ADR['onlem']})*{ADR['kurumsal_pay']}*B5*B2*({ADR['kv']}-B4),0)"),
('KGF maliyeti (TL)',f"=B5*{ADR['kgf_alim']}*{ADR['kgf_pay']}*B2*{ADR['kgf_temerrut']}*(1-{ADR['kgf_kurtarma']})"),
('BT yıllık işletme (TL)',f"={ADR['bt_yillik']}"),('İtiraz gideri (TL)',f"=B5*{ADR['itiraz_oran']}*{ADR['itiraz_maliyet']}"),('Denetim (TL)',f"={ADR['denetim']}*B6"),
('Gelire bağlı yeni prim (TL)',f"=B5*(1-{u})*{E('O')}"),
('NET HAZİNE ETKİSİ (TL/yıl)',"=B6+B17-B7-B8-B9-B10-B11-B12-B13-B14-B15-B16"),('Net (milyar TL)',"=B18/1000000000")]
for i,(a,f) in enumerate(SR): s.cell(row=2+i,column=1,value=a); s.cell(row=2+i,column=2,value=f)
s['A18'].font=b; s['B18'].font=b; s.column_dimensions['A'].width=62; s.column_dimensions['B'].width=22
s['A22']='Not: eski sabit prim kalktı; yerine gelire bağlı prim (Girdiler: prim_k, prim_ay); varyantlar Senaryolar_Python sayfasındadır.'
# --- Python sonuçları (statik)
p=wb.create_sheet('Senaryolar_Python'); p['A1']='Python modeli çıktıları (statik değerler; model.py ve analiz.py ile üretilir)'; p['A1'].font=b
r=3
for j,h in enumerate(['Senaryo','4. yıl net (mlr TL)','5 yıl NPV (mlr TL)']): p.cell(row=r,column=1+j,value=h).font=b
for ad in ['Kötümser','Orta','İyimser']:
    r+=1; p.cell(row=r,column=1,value=ad); p.cell(row=r,column=2,value=S['senaryo'][ad]['yil4']['net']/1e9); p.cell(row=r,column=3,value=S['senaryo'][ad]['npv5']/1e9)
r+=2; p.cell(row=r,column=1,value='Politika varyantları (Orta, 4. yıl net, mlr TL)').font=b
for k,v in S['varyant'].items(): r+=1; p.cell(row=r,column=1,value=k); p.cell(row=r,column=2,value=v/1e9)
r+=2; p.cell(row=r,column=1,value='Kırılma noktası s* (Orta)').font=b; p.cell(row=r,column=2,value=('yok' if S['s_kirilma'] is None else S['s_kirilma']))
r+=2; p.cell(row=r,column=1,value='Monte Carlo (4000 deneme), 4. yıl net (mlr TL)').font=b
for k,v in S['mc']['net_q'].items(): r+=1; p.cell(row=r,column=1,value='P'+k); p.cell(row=r,column=2,value=v/1e9)
r+=1; p.cell(row=r,column=1,value='P(net>0)'); p.cell(row=r,column=2,value=S['mc']['p_net_pozitif'])
r+=2; p.cell(row=r,column=1,value='Tornado (Orta tabanı, 4. yıl net değişimi, mlr TL)').font=b
for t in S['tornado']: r+=1; p.cell(row=r,column=1,value=t[0]); p.cell(row=r,column=2,value=t[3]/1e9); p.cell(row=r,column=3,value=t[4]/1e9)
p.column_dimensions['A'].width=60
wb.save('MIKRO_MUKELLEF_MALI_MODEL.xlsx'); print('kaydedildi')
